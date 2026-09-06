import hashlib
import hmac
import json
import tempfile
import time
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from rest_framework.test import APIClient
from docx import Document

from .models import Application, Profile, User
from .resume_parser import detect_skills, extract_email, extract_full_name, extract_phone, extract_telegram
from .services import get_profile, infer_target_role, normalize_resume_scoring_profile


class AccountApiTests(TestCase):
    def test_target_role_inference_does_not_default_unrelated_resume_to_java(self):
        self.assertEqual(infer_target_role({"position": "Frontend Engineer", "skills": [{"name": "React"}]}), "UNKNOWN")
        self.assertEqual(infer_target_role({"position": "Java Backend Developer", "skills": []}), "JAVA_BACKEND")
        self.assertEqual(infer_target_role({"position": "Platform Engineer", "skills": [{"name": "Linux"}, {"name": "Kubernetes"}]}), "DEVOPS")

    def test_resume_scoring_profile_preserves_fractional_experience(self):
        normalized, changed = normalize_resume_scoring_profile({"position": "Frontend Engineer", "experience": "1,5 года", "skills": []})
        self.assertTrue(changed)
        self.assertEqual(normalized["targetRole"], "UNKNOWN")
        self.assertEqual(normalized["experienceYears"], 1.5)

    def test_resume_contact_extraction_returns_form_values(self):
        text = "Иванов Иван Иванович\nEmail: ivan@example.com\nТелефон: +7 (999) 123-45-67\nTelegram: @ivan_dev"

        self.assertEqual(extract_full_name(text), "Иванов Иван Иванович")
        self.assertEqual(extract_email(text), "ivan@example.com")
        self.assertEqual(extract_phone(text), "+7 (999) 123-45-67")
        self.assertEqual(extract_telegram(text), "@ivan_dev")

    def test_resume_contact_extraction_uses_filename_when_pdf_text_is_corrupted(self):
        self.assertEqual(
            extract_full_name("�������� ������\nDevOps �������", "Резюме_DevOps_инженер_Матвеи_Пасечник_от_04.pdf"),
            "Матвеи Пасечник",
        )

    @override_settings(LOCAL_AUTH_BYPASS=True, DEBUG=True)
    def test_local_auth_bypass_uses_shared_account(self):
        first_client = APIClient(enforce_csrf_checks=True)
        second_client = APIClient()

        first = first_client.get("/api/auth/me/")
        second = second_client.get("/api/auth/me/")

        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(first.data["user"]["email"], "local@jobs.dev")
        self.assertEqual(first.data["user"]["id"], second.data["user"]["id"])
        self.assertEqual(User.objects.filter(email="local@jobs.dev").count(), 1)
        self.assertTrue(Profile.objects.filter(user__email="local@jobs.dev").exists())

        write = first_client.post("/api/profile/saved/", {"jobId": 123, "saved": True}, format="json")
        self.assertEqual(write.status_code, 200)
        self.assertEqual(write.data["savedJobIds"], [123])

    @override_settings(LOCAL_AUTH_BYPASS=True, DEBUG=True)
    def test_extension_archive_can_be_downloaded_locally(self):
        response = APIClient().get("/api/extension/download/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        self.assertIn("jobs-dev-zen-extension.zip", response["Content-Disposition"])

    @override_settings(LOCAL_AUTH_BYPASS=True, DEBUG=True)
    def test_chrome_extension_archive_can_be_downloaded_locally(self):
        response = APIClient().get("/api/extension/download/chrome/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        self.assertIn("jobs-dev-zen-extension-chrome.zip", response["Content-Disposition"])

    def test_extension_submit_is_resolved_to_catalog_application(self):
        with tempfile.NamedTemporaryFile(mode="w", suffix=".json", encoding="utf-8", delete=False) as catalog:
            json.dump({"vacancies": [{"id": "42", "source_key": "acme", "company": "Acme", "title": "Senior Python Engineer", "location": "Remote", "url": "https://acme.example/jobs/python-42", "technologies": ["Python"]}]}, catalog)
            catalog_path = catalog.name
        try:
            with override_settings(VACANCY_CATALOG_PATH=catalog_path, JOB_TRACKER_DB_PATH=""):
                client = APIClient()
                client.post("/api/auth/email/", {"email": "extension@example.com", "password": "correct-horse", "mode": "register"}, format="json")
                response = client.post("/api/applications/from-extension/", {"pageUrl": "https://acme.example/jobs/python-42?utm_source=board", "title": "Senior Python Engineer", "company": "Acme"}, format="json")
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.data["matched"])
            self.assertEqual(response.data["match"], "url")
            self.assertEqual(response.data["application"]["detectedBy"], "extension")
            self.assertEqual(Application.objects.get(user__email="extension@example.com").payload["externalId"], "42")
        finally:
            import os
            os.unlink(catalog_path)

    def test_extension_submit_does_not_create_false_positive(self):
        with tempfile.NamedTemporaryFile(mode="w", suffix=".json", encoding="utf-8", delete=False) as catalog:
            json.dump({"vacancies": [{"id": "42", "source_key": "acme", "company": "Acme", "title": "Senior Python Engineer", "url": "https://acme.example/jobs/python-42"}]}, catalog)
            catalog_path = catalog.name
        try:
            with override_settings(VACANCY_CATALOG_PATH=catalog_path, JOB_TRACKER_DB_PATH=""):
                client = APIClient()
                client.post("/api/auth/email/", {"email": "extension-miss@example.com", "password": "correct-horse", "mode": "register"}, format="json")
                response = client.post("/api/applications/from-extension/", {"pageUrl": "https://other.example/jobs/unrelated", "title": "Accountant", "company": "Other"}, format="json")
            self.assertEqual(response.status_code, 200)
            self.assertFalse(response.data["matched"])
            self.assertFalse(Application.objects.filter(user__email="extension-miss@example.com").exists())
        finally:
            import os
            os.unlink(catalog_path)

    def test_devops_scoring_uses_confirmed_resume_skills(self):
        client = APIClient()
        client.post("/api/auth/email/", {"email": "devops@example.com", "password": "correct-horse", "mode": "register"}, format="json")
        profile = Profile.objects.get(user__email="devops@example.com")
        profile.resume = {
            "position": "DevOps engineer", "targetRole": "DEVOPS",
            "skills": [{"name": "Docker", "confirmed": True}, {"name": "Kubernetes", "confirmed": True}],
        }
        profile.save()

        response = client.post("/api/scoring/rank/", {"items": [{"id": 7, "features": {"v": "1.0.0", "p": "DEVOPS", "r": ["DEVOPS", 1, 1], "s": ["MIDDLE", 1], "e": 3, "c": [["DOCKER", 1, 1, 1, 1], ["KUBERNETES", 1, 1, 1, 1], ["TERRAFORM", 1, 1, 1, 1]], "n": [], "d": ""}}]}, format="json")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["profile"]["targetProfile"], "DEVOPS")
        self.assertEqual(response.data["scores"][0]["vacancyId"], "catalog:7")
        self.assertEqual(response.data["scores"][0]["scoringVersion"], "2.0.0")
        self.assertIn(response.data["scores"][0]["eligibility"], {"ELIGIBLE", "INELIGIBLE", "UNCERTAIN"})
        self.assertGreaterEqual(response.data["scores"][0]["confidence"], 0)
        self.assertGreater(response.data["scores"][0]["score"], 0)

    @override_settings(TELEGRAM_BOT_TOKEN="test-token", TELEGRAM_BOT_USERNAME="jobsdev_bot")
    def test_telegram_login_creates_session_and_profile(self):
        payload = {"id": "42", "first_name": "Ada", "username": "ada", "auth_date": str(int(time.time()))}
        check_string = "\n".join(f"{key}={payload[key]}" for key in sorted(payload))
        secret_key = hashlib.sha256(b"test-token").digest()
        payload["hash"] = hmac.new(secret_key, check_string.encode(), hashlib.sha256).hexdigest()

        response = APIClient().post("/api/auth/telegram/", payload, format="json")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["user"]["telegram"], "@ada")
        self.assertTrue(User.objects.filter(telegram_id="42").exists())
        self.assertTrue(Profile.objects.filter(user__telegram_id="42").exists())

    @override_settings(TELEGRAM_SYNC_TOKEN="sync-secret")
    def test_telegram_subscribers_return_only_opted_in_users(self):
        user = User.objects.create_user(email="telegram@example.com", password="correct-horse", telegram_id="777")
        profile = get_profile(user)
        profile.settings["notifications"].update({
            "telegramEnabled": True,
            "newJobs": True,
            "telegramKeywords": ["backend"],
        })
        profile.onboarding = {"roles": ["backend"], "levels": [], "formats": []}
        profile.save(update_fields=["settings", "onboarding", "updated_at"])
        response = APIClient().get("/api/integrations/telegram/subscribers/", HTTP_AUTHORIZATION="Bearer sync-secret")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["subscribers"][0]["chatId"], "777")
        self.assertEqual(response.data["subscribers"][0]["filter"]["keywords"], ["backend"])

    def test_email_registration_and_profile_patch_persist(self):
        client = APIClient()
        response = client.post("/api/auth/email/", {"email": "ada@example.com", "password": "correct-horse", "mode": "register"}, format="json")
        self.assertEqual(response.status_code, 200)
        response = client.patch("/api/profile/", {"name": "Ada Lovelace", "onboarding": {"roles": ["backend"]}}, format="json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["user"]["name"], "Ada Lovelace")
        self.assertEqual(Profile.objects.get(user__email="ada@example.com").onboarding, {"roles": ["backend"]})

    def test_resume_upload_extracts_docx_profile_and_skills(self):
        client = APIClient()
        client.post("/api/auth/email/", {"email": "resume@example.com", "password": "correct-horse", "mode": "register"}, format="json")
        document = Document()
        document.add_paragraph("Иванов Иван Иванович")
        document.add_paragraph("Senior Backend Engineer")
        document.add_paragraph("Опыт: 6 лет. Python, Golang, PostgreSQL, Docker, Kubernetes, REST API")
        document.add_paragraph("Email: ivan@example.com | Телефон: +7 (999) 123-45-67 | Telegram: @ivan_dev")
        file_data = BytesIO()
        document.save(file_data)

        response = client.post(
            "/api/profile/resume/",
            {"resume": SimpleUploadedFile("resume.docx", file_data.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            format="multipart",
        )

        self.assertEqual(response.status_code, 200)
        resume = response.data["resume"]
        self.assertEqual(resume["position"], "Senior Backend Engineer")
        self.assertEqual(resume["experience"], "6 лет")
        self.assertEqual(resume["fullName"], "Иванов Иван Иванович")
        self.assertEqual(resume["contactEmail"], "ivan@example.com")
        self.assertEqual(resume["contactPhone"], "+7 (999) 123-45-67")
        self.assertEqual(resume["contactTelegram"], "@ivan_dev")
        self.assertTrue(resume["id"].startswith("upload-"))
        self.assertTrue(resume["hasFile"])
        self.assertEqual({skill["name"] for skill in resume["skills"]}, {"Python", "Go", "PostgreSQL", "Docker", "Kubernetes", "REST API"})

        downloaded = client.get("/api/profile/resume/file/")
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(b"".join(downloaded.streaming_content), file_data.getvalue())

    def test_resume_upload_rejects_unsupported_format(self):
        client = APIClient()
        client.post("/api/auth/email/", {"email": "resume-format@example.com", "password": "correct-horse", "mode": "register"}, format="json")

        response = client.post(
            "/api/profile/resume/",
            {"resume": SimpleUploadedFile("resume.txt", b"Python", content_type="text/plain")},
            format="multipart",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("PDF и DOCX", response.data["message"])

    def test_multiple_resume_uploads_are_kept_and_can_be_selected(self):
        client = APIClient()
        client.post("/api/auth/email/", {"email": "resume-list@example.com", "password": "correct-horse", "mode": "register"}, format="json")

        def make_resume(filename: str, title: str) -> bytes:
            document = Document()
            document.add_paragraph("Тестовый кандидат")
            document.add_paragraph(title)
            document.add_paragraph("Опыт: 3 года. Python, Docker")
            output = BytesIO()
            document.save(output)
            return output.getvalue()

        first_data = make_resume("first.docx", "Backend Engineer")
        second_data = make_resume("second.docx", "DevOps Engineer")
        first = client.post("/api/profile/resume/", {"resume": SimpleUploadedFile("first.docx", first_data)}, format="multipart")
        second = client.post("/api/profile/resume/", {"resume": SimpleUploadedFile("second.docx", second_data)}, format="multipart")
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(len(second.data["resumes"]), 2)
        self.assertNotEqual(first.data["resume"]["id"], second.data["resume"]["id"])
        self.assertTrue(second.data["resume"]["isActive"])

        selected = client.patch("/api/profile/resume/", {"activeResumeId": first.data["resume"]["id"]}, format="json")
        self.assertEqual(selected.status_code, 200)
        self.assertEqual(selected.data["resume"]["id"], first.data["resume"]["id"])
        downloaded = client.get(f"/api/profile/resume/file/?id={first.data['resume']['id']}")
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(b"".join(downloaded.streaming_content), first_data)

    def test_skill_detection_handles_aliases_and_pdf_line_breaks(self):
        skills = detect_skills("Golang, K8s, Postgre-\nsql, JavaScript, Google Cloud")

        self.assertEqual({skill["name"] for skill in skills}, {"Go", "Kubernetes", "PostgreSQL", "JavaScript", "GCP"})

    def test_experience_ignores_age_and_prefers_work_experience(self):
        from .resume_parser import extract_experience, extract_experience_duration

        self.assertEqual(extract_experience("Мужчина, 23 года\nОпыт работы — 4 года 3 месяца"), "4 года")
        self.assertEqual(extract_experience_duration("Мужчина, 23 года\nОпыт работы — 4 года 3 месяца"), (4.25, 3))

    def test_old_resume_analysis_is_refreshed_when_account_is_loaded(self):
        from django.core.files.base import ContentFile
        from .services import account_payload, get_profile

        user = User.objects.create_user(email="refresh@example.com", password="correct-horse")
        profile = get_profile(user)
        document = Document()
        document.add_paragraph("Java Backend Developer")
        document.add_paragraph("Опыт работы — 4 года. Java, Spring Boot, PostgreSQL")
        file_data = BytesIO()
        document.save(file_data)

        profile.resume_file.save("resume.docx", ContentFile(file_data.getvalue()), save=False)
        profile.resume = {"fileName": "resume.docx", "experience": "", "position": "", "skills": []}
        profile.save()

        payload = account_payload(user)

        self.assertEqual(payload["resume"]["analysisVersion"], 4)
        self.assertEqual(payload["resume"]["position"], "Java Backend Developer")
        self.assertEqual(payload["resume"]["experience"], "4 года")
        self.assertEqual({skill["name"] for skill in payload["resume"]["skills"]}, {"Java", "Spring", "Spring Boot", "PostgreSQL"})
